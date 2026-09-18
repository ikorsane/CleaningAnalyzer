import cv2
import numpy as np
import pandas as pd
from pathlib import Path

BOTTOM_MARGIN_FRAC = 0.05

# ------------------------- geometry -------------------------

def order4(pts):
    s=pts.sum(axis=1); d=np.diff(pts,axis=1).ravel()
    return np.array([pts[np.argmin(s)], pts[np.argmin(d)],
                     pts[np.argmax(s)], pts[np.argmax(d)]],np.float32)

def rectify(img, pts):
    tl,tr,br,bl=order4(pts)
    W=int(max(np.linalg.norm(br-bl),np.linalg.norm(tr-tl)))
    H=int(max(np.linalg.norm(tr-br),np.linalg.norm(tl-bl)))
    dst=np.float32([[0,0],[W-1,0],[W-1,H-1],[0,H-1]])
    M=cv2.getPerspectiveTransform(np.float32([tl,tr,br,bl]),dst)
    return cv2.warpPerspective(img,M,(W,H))

# ------------------------- color / baseline -------------------------

def lab_float(img):
    return cv2.cvtColor(img,cv2.COLOR_BGR2LAB).astype(np.float32)

def roi_mask(shape, r):
    x,y,w,h=r
    m=np.zeros(shape[:2],np.uint8); m[y:y+h,x:x+w]=255
    return m

def robust_lab(img_lab, mask):
    p=img_lab[mask>0]
    return np.median(p,axis=0)

def spatial_dirty_model(lab, excluded, sigma_frac=.10):
    """
    Estimate the local dirty appearance from untouched pixels.
    Normalized convolution prevents excluded/cleaned pixels from bleeding in.
    """
    h,w=lab.shape[:2]
    valid=(excluded==0).astype(np.float32)
    sigma=max(15.0, sigma_frac*min(h,w))
    den=cv2.GaussianBlur(valid,(0,0),sigmaX=sigma,sigmaY=sigma)
    out=np.zeros_like(lab,np.float32)
    for c in range(3):
        num=cv2.GaussianBlur(lab[:,:,c]*valid,(0,0),sigmaX=sigma,sigmaY=sigma)
        out[:,:,c]=num/np.maximum(den,1e-4)
    # Fill poorly supported pixels with global dirty median.
    g=np.median(lab[valid>0.5],axis=0)
    bad=den<0.05
    out[bad]=g
    return out

def cleaning_fraction(lab, dirty_model, clean_lab):
    """
    Project each pixel from its local dirty endpoint toward the clean endpoint
    in 3D CIELAB. 0 = dirty, 1 = clean. This uses color direction, not whiteness.
    """
    v=clean_lab.reshape(1,1,3)-dirty_model
    q=lab-dirty_model
    den=np.sum(v*v,axis=2)
    f=np.sum(q*v,axis=2)/np.maximum(den,1e-6)
    return np.clip(f,0,1)

def originally_soiled_mask(dirty_model, clean_lab, clean_distance_fraction=0.45):
    """
    Estimate substrate that was originally covered by soil.
    Uses the estimated pre-cleaning dirty endpoint, not the observed cleaned colour.
    """
    clean=clean_lab.reshape(1,1,3)
    d=np.linalg.norm(dirty_model-clean,axis=2)
    positive=d[d>1e-6]
    if positive.size==0:
        return np.ones(d.shape,np.uint8)*255
    dirty_sep=float(np.percentile(positive,75))
    threshold=clean_distance_fraction*dirty_sep
    m=(d>=threshold).astype(np.uint8)*255
    k=max(3,int(0.006*min(d.shape))); k += 1-k%2
    kernel=cv2.getStructuringElement(cv2.MORPH_ELLIPSE,(k,k))
    m=cv2.morphologyEx(m,cv2.MORPH_OPEN,kernel,iterations=1)
    m=cv2.morphologyEx(m,cv2.MORPH_CLOSE,kernel,iterations=1)
    return m

# ------------------------- footprint proposal -------------------------

def regularize_tongue(mask, guide):
    """
    Preserve a rounded/irregular detected top. In the lower part, once a stable
    footprint width is established, carry that width vertically to plate bottom.
    """
    h,w=mask.shape
    ys,xs=np.where(mask>0)
    if len(xs)==0: return mask
    y0,y1=ys.min(),ys.max()
    out=np.zeros_like(mask)

    # Row envelopes, tolerant of ragged/partial cleaning.
    rows=[]
    for y in range(y0,y1+1):
        xx=np.where(mask[y]>0)[0]
        if len(xx)>=3:
            rows.append((y,np.percentile(xx,5),np.percentile(xx,95)))
    if not rows: return mask
    arr=np.array(rows,float)
    # Smooth side envelopes.
    L=cv2.GaussianBlur(arr[:,1].astype(np.float32).reshape(-1,1),(1,0),sigmaX=0,sigmaY=8).ravel()
    R=cv2.GaussianBlur(arr[:,2].astype(np.float32).reshape(-1,1),(1,0),sigmaX=0,sigmaY=8).ravel()
    yy=arr[:,0].astype(int)
    for y,l,r in zip(yy,L,R):
        out[y,max(0,int(l)):min(w,int(r)+1)]=255

    # Carry established lower width to bottom.
    tail_start=int(y0+0.62*(y1-y0))
    band=(yy>=tail_start)
    if np.any(band):
        l=int(np.median(L[band])); r=int(np.median(R[band]))
        # Avoid narrowing: include existing envelope too.
        for y in range(tail_start,h):
            xx=np.where(out[y]>0)[0]
            if len(xx):
                l2=min(l,xx.min()); r2=max(r,xx.max())
            else:
                l2,r2=l,r
            out[y,max(0,l2):min(w,r2+1)]=255

    # Close modest gaps but retain overall footprint shape.
    k=max(5,int(0.012*w)); k += 1-k%2
    out=cv2.morphologyEx(out,cv2.MORPH_CLOSE,
                         cv2.getStructuringElement(cv2.MORPH_ELLIPSE,(k,k)),iterations=1)
    return out

def propose_footprint(clean_frac, guide_roi):
    """
    Propose the contacted product track from a rough user guide.

    Important: the guide is not a hard top/bottom crop, but it DOES define the
    track's horizontal corridor.  This prevents an upward search for a rounded
    tip from wandering into handwriting, plate edges, or neighbouring tracks.
    """
    x,y,w,h=guide_roi
    H,W=clean_frac.shape

    # Keep the search tightly tied to the product's horizontal position.
    # Only a small side allowance is permitted; the useful extra search is
    # mainly ABOVE the guide so a rounded tip cannot be clipped flat.
    side=max(4,int(0.04*w))
    top=max(12,int(0.22*h))
    bottom=max(4,int(0.03*h))
    x0=max(0,x-side); x1=min(W,x+w+side)
    y0=max(0,y-top);  y1=min(H,y+h+bottom)
    local=clean_frac[y0:y1,x0:x1]

    seed=(local>0.08).astype(np.uint8)*255
    seed=cv2.medianBlur(seed,5)

    # Use only modest closing.  The previous large kernel could physically
    # bridge the product tip to bright handwriting above it.
    k=max(3,int(0.012*min(w,h))); k += 1-k%2
    seed=cv2.morphologyEx(seed,cv2.MORPH_CLOSE,
                         cv2.getStructuringElement(cv2.MORPH_ELLIPSE,(k,k)),
                         iterations=1)

    n,labels,stats,cent=cv2.connectedComponentsWithStats(seed)
    if n<=1:
        m=np.zeros(clean_frac.shape,np.uint8)
        m[y:y+h,x:x+w]=255
        return m

    # Pick the component with the strongest overlap with the ORIGINAL guide.
    gx0=x-x0; gx1=gx0+w
    gy0=y-y0; gy1=gy0+h
    best=None; best_score=-1
    for i in range(1,n):
        comp=(labels==i)
        overlap=np.count_nonzero(comp[gy0:gy1,gx0:gx1])
        if overlap==0:
            continue
        area=stats[i,cv2.CC_STAT_AREA]
        score=overlap + 0.005*area
        if score>best_score:
            best_score=score; best=i

    if best is None:
        m=np.zeros(clean_frac.shape,np.uint8)
        m[y:y+h,x:x+w]=255
        return m

    localmask=(labels==best).astype(np.uint8)*255

    # Reject implausible sideways excursions.  A product track may taper toward
    # its tip, but it should not suddenly become much wider than the rough track.
    centre=x + 0.5*w
    max_half=0.60*w
    corridor=np.zeros_like(localmask)
    global_left=max(x0, int(centre-max_half))
    global_right=min(x1, int(centre+max_half))
    corridor[:, global_left-x0:global_right-x0]=255
    localmask=cv2.bitwise_and(localmask,corridor)

    # Keep only the connected piece that still overlaps the guide after the
    # corridor restriction.  This removes detached lettering/edge fragments.
    n2,lab2,stats2,_=cv2.connectedComponentsWithStats(localmask)
    if n2>1:
        best2=None; score2=-1
        for i in range(1,n2):
            comp=(lab2==i)
            ov=np.count_nonzero(comp[gy0:gy1,gx0:gx1])
            if ov>score2:
                score2=ov; best2=i
        if best2 is not None:
            localmask=(lab2==best2).astype(np.uint8)*255

    full=np.zeros(clean_frac.shape,np.uint8)
    full[y0:y1,x0:x1]=localmask
    return regularize_tongue(full,guide_roi)


# ------------------------- analysis -------------------------

def heatmap_overlay(img, frac, mask, name):
    heat=(np.clip(frac*255,0,255)).astype(np.uint8)
    heat=cv2.applyColorMap(heat,cv2.COLORMAP_TURBO)
    out=img.copy()
    alpha=.55
    ix=mask>0
    out[ix]=(alpha*heat[ix]+(1-alpha)*out[ix]).astype(np.uint8)
    cnts,_=cv2.findContours(mask,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)
    cv2.drawContours(out,cnts,-1,(0,0,255),3)
    cv2.putText(out,"Blue = low cleaning   Red = high cleaning",
                (20,35),cv2.FONT_HERSHEY_SIMPLEX,.8,(255,255,255),2,cv2.LINE_AA)
    return out

def bottom_exclusion_mask(shape, frac=BOTTOM_MARGIN_FRAC):
    """Mask of the valid plate area after removing a fixed bottom margin."""
    h, w = shape[:2]
    m = np.ones((h, w), np.uint8) * 255
    cut = int(round(h * (1.0 - float(frac))))
    cut = max(0, min(h, cut))
    m[cut:, :] = 0
    return m

def metrics(name, frac, mask):
    vals=frac[mask>0]
    area=int(np.count_nonzero(mask))
    return {
        "Product":name,
        "Footprint area (px)":area,
        "Mean cleaning depth (%)":100*float(np.mean(vals)),
        "Median cleaning depth (%)":100*float(np.median(vals)),
        "25th percentile cleaning (%)":100*float(np.percentile(vals,25)),
        "75th percentile cleaning (%)":100*float(np.percentile(vals,75)),
        "Footprint >=50% cleaned (%)":100*float(np.mean(vals>=.50)),
        "Footprint >=75% cleaned (%)":100*float(np.mean(vals>=.75)),
        "Integrated optical removal (equiv clean px)":float(np.sum(vals)),
    }


def save_report_workbook(df, outpath):
    """Create report workbook with replicate-level results and across-replicate mean/SD."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.utils import get_column_letter

    metric_map = {
        "Cleaning performance (%)": "Mean cleaning depth (%)",
        "Area >=50% cleaned (%)": "Footprint >=50% cleaned (%)",
        "Area >=75% cleaned (%)": "Footprint >=75% cleaned (%)",
        "Footprint area (px)": "Footprint area (px)",
    }
    g=df.groupby("Product", sort=False)
    rows=[]
    for product, sub in g:
        row={"Product":product,"n":len(sub)}
        for label,col in metric_map.items():
            row[label]=sub[col].mean()
            row[label+" SD"]=sub[col].std(ddof=1) if len(sub)>1 else np.nan
        rows.append(row)
    summary=pd.DataFrame(rows)
    summary["Rank"]=summary["Cleaning performance (%)"].rank(ascending=False,method="min").astype(int)

    wb=Workbook(); ws=wb.active; ws.title="Summary"
    rep=wb.create_sheet("Replicate Results"); detail=wb.create_sheet("Detailed Results")
    navy="17365D"; blue="4472C4"; pale="D9EAF7"; white="FFFFFF"; thin=Side(style="thin",color="D9D9D9")

    ws["A1"]="Cleaning Performance Summary"; ws["A1"].font=Font(bold=True,size=18,color=navy)
    ws.merge_cells("A1:K1")
    ws["A2"]="Mean ± SD across independent plate/image replicates. Each replicate is weighted equally."; ws["A2"].font=Font(size=10,color="666666")
    ws.merge_cells("A2:K2")
    cols=["Product","n","Cleaning performance (%)","Cleaning performance (%) SD","Area >=50% cleaned (%)","Area >=50% cleaned (%) SD","Area >=75% cleaned (%)","Area >=75% cleaned (%) SD","Footprint area (px)","Footprint area (px) SD","Rank"]
    ws["A4"]="Comparative results"; ws["A4"].font=Font(bold=True,color=white); ws["A4"].fill=PatternFill("solid",fgColor=navy); ws.merge_cells("A4:K4")
    for c,col in enumerate(cols,1):
        cell=ws.cell(5,c,col); cell.font=Font(bold=True,color=white); cell.fill=PatternFill("solid",fgColor=blue); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for r,row in enumerate(summary[cols].itertuples(index=False,name=None),6):
        for c,val in enumerate(row,1):
            if pd.isna(val): val=None
            cell=ws.cell(r,c,val.item() if hasattr(val,"item") else val); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin); cell.alignment=Alignment(horizontal="left" if c==1 else "center")
            if c in (3,4,5,6,7,8): cell.number_format="0.0"
            elif c in (2,9,10,11): cell.number_format="0"
    last=5+len(summary)
    for col in ("C","E","G"):
        ws.conditional_formatting.add(f"{col}6:{col}{last}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=100,color=pale,showValue=True))

    # Clean report-style charts. SD is shown numerically in the adjacent summary columns.
    ch=BarChart(); ch.type="col"; ch.style=10; ch.title="Mean cleaning performance"; ch.y_axis.title="Cleaning (%)"; ch.y_axis.scaling.min=0; ch.y_axis.scaling.max=100; ch.height=8; ch.width=14; ch.add_data(Reference(ws,min_col=3,min_row=5,max_row=last),titles_from_data=True); ch.set_categories(Reference(ws,min_col=1,min_row=6,max_row=last)); ch.legend=None; ch.dLbls=DataLabelList(); ch.dLbls.showVal=True; ws.add_chart(ch,"M2")
    ch2=BarChart(); ch2.type="col"; ch2.style=10; ch2.title="Cleaning coverage"; ch2.y_axis.title="Footprint meeting threshold (%)"; ch2.y_axis.scaling.min=0; ch2.y_axis.scaling.max=100; ch2.height=8; ch2.width=14; ch2.add_data(Reference(ws,min_col=5,max_col=7,min_row=5,max_row=last),titles_from_data=True,from_rows=False); ch2.series=[ch2.series[0],ch2.series[2]] if len(ch2.series)>=3 else ch2.series; ch2.set_categories(Reference(ws,min_col=1,min_row=6,max_row=last)); ch2.legend.position="b"; ws.add_chart(ch2,"U2")
    note=ws.cell(last+3,1); ws.merge_cells(start_row=last+3,start_column=1,end_row=last+6,end_column=11); note.value="Interpretation: values are arithmetic means across replicate plates/images; SD is the sample standard deviation between replicates (n=1: SD not reported). Replicates are weighted equally rather than by footprint pixel count. Footprint area is a diagnostic of contact/wetting and should not by itself be interpreted as cleaning efficacy."; note.font=Font(size=9,color="666666"); note.alignment=Alignment(wrap_text=True,vertical="top")
    for c,w in {"A":18,"B":8,"C":22,"D":18,"E":22,"F":18,"G":22,"H":18,"I":20,"J":18,"K":9}.items(): ws.column_dimensions[c].width=w
    ws.row_dimensions[5].height=42; ws.sheet_view.showGridLines=False

    # Replicate-level sheet: one row per product per plate.
    repdf=df.copy()
    for c,col in enumerate(repdf.columns,1):
        cell=rep.cell(1,c,col); cell.font=Font(bold=True,color=white); cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(horizontal="center",wrap_text=True)
    for r,vals in enumerate(repdf.itertuples(index=False,name=None),2):
        for c,val in enumerate(vals,1): rep.cell(r,c,val.item() if hasattr(val,"item") else val).number_format="0.00" if c>2 else "General"
    rep.freeze_panes="C2"; rep.auto_filter.ref=rep.dimensions; rep.sheet_view.showGridLines=False

    # Detailed results currently equal replicate results, retained for compatibility/report traceability.
    for c,col in enumerate(df.columns,1):
        cell=detail.cell(1,c,col); cell.font=Font(bold=True,color=white); cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(horizontal="center",wrap_text=True)
    for r,vals in enumerate(df.itertuples(index=False,name=None),2):
        for c,val in enumerate(vals,1): detail.cell(r,c,val.item() if hasattr(val,"item") else val).number_format="0.00" if c>2 else "General"
    for sh in (rep,detail):
        for c in range(1,len(df.columns)+1): sh.column_dimensions[get_column_letter(c)].width=18 if c<=2 else 24
        sh.freeze_panes="C2"; sh.auto_filter.ref=sh.dimensions; sh.sheet_view.showGridLines=False
    wb.save(outpath)

